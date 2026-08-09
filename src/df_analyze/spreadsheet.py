from __future__ import annotations

import csv
import shlex
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


def _option_line(values: Iterable[object]) -> str | None:
    cells = [str(value).strip() for value in values if str(value).strip()]
    if not cells or not cells[0].startswith("--"):
        return None
    return shlex.join(cells)


def _csv_options(path: Path, separator: str) -> str:
    options: list[str] = []
    with path.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.reader(handle, delimiter=separator):
            if not any(str(value).strip() for value in row):
                continue
            line = _option_line(row)
            if line is None:
                break
            options.append(line)
    return " ".join(options)


def _xlsx_options(path: Path) -> str:
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if len(workbook.sheetnames) != 1:
            raise RuntimeError(
                "Found multiple sheets in Excel workbook. Make sure the Excel file "
                "has only a single sheet formatted correctly for df-analyze."
            )
        worksheet = workbook[workbook.sheetnames[0]]
        options: list[str] = []
        for row in worksheet.iter_rows(values_only=True):
            if not any(value is not None and str(value).strip() for value in row):
                continue
            line = _option_line(value for value in row if value is not None)
            if line is None:
                break
            options.append(line)
        return " ".join(options)
    finally:
        workbook.close()


def spreadsheet_options(path: Path, separator: str = ",") -> str:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _csv_options(path, separator)
    if suffix == ".xlsx":
        return _xlsx_options(path)
    raise ValueError(f"Unsupported spreadsheet extension: {path.suffix}")
